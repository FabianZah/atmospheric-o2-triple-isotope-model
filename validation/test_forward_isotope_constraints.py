"""Forward propagation accuracy, domain, API, and export contracts."""

from io import BytesIO
from itertools import product
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path
from threading import Barrier, Lock
import subprocess
import sys
import time

import numpy as np
import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from scipy.special import ndtr, ndtri
from scipy.optimize import brentq

import forward_isotope_constraints as propagation
from forward_isotope_constraints import ForwardIsotopeConstraints, predict_isotopes
from updated_constrained_pco2_posterior import CoordinateConstraint as C
from updated_output_surface import UpdatedOutputSurfaceInput, load_updated_output_surface
import web_api


class LinearSurface:
    domain = {"pco2_ppm": (50., 60000.), "gpp_pgC_per_year": (18.256264, 850.), "po2_pal": (.1, 2.)}

    def evaluate_central_cap_delta17_grid(self, *, p_o2_pal, p_co2_ppm, gpp_pgC_per_year):
        return -.001 * p_co2_ppm + .01 * gpp_pgC_per_year - p_o2_pal

    def evaluate_central_delta18_prime_grid(self, **args):
        conventional = 24. + .1 * self.evaluate_central_cap_delta17_grid(**args)
        return 1000. * np.log1p(conventional / 1000.)


def test_fixed_matches_existing_forward():
    actual = predict_isotopes(ForwardIsotopeConstraints())
    reference = load_updated_output_surface().evaluate(UpdatedOutputSurfaceInput())
    assert actual["isotopes"]["cap_delta17_prime_permil"]["median"] == pytest.approx(reference.central_cap_delta17_prime_permil, abs=1e-12)
    assert actual["isotopes"]["delta18_conventional_permil"]["median"] == pytest.approx(1000 * np.expm1(reference.central_delta18_prime_permil / 1000), abs=1e-12)
    assert actual["isotopes"]["cap_delta17_prime_permil"]["interval95"] is None


@pytest.mark.parametrize("constraint", [C("range", lower=100, upper=500), C("normal", center=80, sigma=60)])
def test_univariate_analytic_quantiles_and_delta18_conversion(constraint):
    request = ForwardIsotopeConstraints(pco2_constraint=constraint)
    actual = predict_isotopes(request, surface=LinearSurface())
    bounds = actual["effective_bounds"]["pCO2"]
    probabilities = np.array([.025, .5, .975])
    if constraint.kind == "range":
        co2 = bounds[0] + probabilities * (bounds[1] - bounds[0])
    else:
        lo, hi = ndtr((np.asarray(bounds) - constraint.center) / constraint.sigma)
        co2 = constraint.center + constraint.sigma * ndtri(lo + probabilities * (hi-lo))
    expected = (-.001 * co2 + 1.9)[::-1]
    iso = actual["isotopes"]["cap_delta17_prime_permil"]
    np.testing.assert_allclose([iso["interval95"][0], iso["median"], iso["interval95"][1]], expected, atol=.0001)
    np.testing.assert_allclose(actual["isotopes"]["delta18_conventional_permil"]["interval95"], 24 + .1 * expected[[0, 2]], atol=.0001)
    assert actual == predict_isotopes(request, surface=LinearSurface())


def test_three_uncertainties_have_analytic_mean_and_variance():
    result = predict_isotopes(ForwardIsotopeConstraints(
        C("range", lower=100, upper=500), C("range", lower=200, upper=400), C("range", lower=.5, upper=1.5)), surface=LinearSurface())
    isotope = result["isotopes"]["cap_delta17_prime_permil"]
    assert isotope["mean"] == pytest.approx(1.7, abs=.0001)
    assert isotope["standard_deviation"] == pytest.approx(np.sqrt((.4**2 + 2**2 + 1**2)/12), abs=.0001)
    widths = np.array([.4, 2., 1.])
    def cdf(value):
        return sum((-1)**sum(bits) * max(value - np.dot(bits, widths), 0)**3
                   for bits in product((0, 1), repeat=3)) / (6 * np.prod(widths))
    exact = [brentq(lambda x: cdf(x) - q, 0, 3.4) for q in (.025, .975)]
    np.testing.assert_allclose(isotope["interval95"], exact, atol=.0001 + .0002 * (exact[1] - exact[0]))


@pytest.mark.parametrize('coordinate,coefficient,bounds,center,sigma', [
    ('pco2_constraint', -.001, (100, 500), 80, 60),
    ('gpp_constraint', .01, (100, 500), 40, 30),
    ('po2_constraint', -1., (.3, 1.3), .2, .2),
])
@pytest.mark.parametrize('kind', ['range', 'normal'])
def test_each_input_uncertainty_matches_analytic_marginal(coordinate, coefficient, bounds, center, sigma, kind):
    from dataclasses import replace
    constraint = C('range', lower=bounds[0], upper=bounds[1]) if kind == 'range' else C('normal', center=center, sigma=sigma)
    request = replace(ForwardIsotopeConstraints(), **{coordinate: constraint})
    actual = predict_isotopes(request, surface=LinearSurface())
    name = {'pco2_constraint':'pCO2', 'gpp_constraint':'GPP', 'po2_constraint':'pO2'}[coordinate]
    low, high = actual['effective_bounds'][name]
    q = np.array([.025, .5, .975])
    if kind == 'range':
        values = low + q * (high-low)
    else:
        a, b = ndtr((np.array([low,high])-center)/sigma)
        values = center + sigma * ndtri(a + q*(b-a))
    default = {'pco2_constraint':294., 'gpp_constraint':290., 'po2_constraint':1.}[coordinate]
    default_isotope = -.001 * 294. + .01 * 290. - 1.
    expected = np.sort(default_isotope + coefficient*(values-default))
    isotope = actual['isotopes']['cap_delta17_prime_permil']
    np.testing.assert_allclose([isotope['interval95'][0], isotope['median'], isotope['interval95'][1]], expected, atol=.0001)


def test_sorted_replicate_statistics_match_concatenated_numpy():
    rng = np.random.default_rng(7)
    for size in (2, 10, 321):
        arrays = rng.normal(size=(2, 2, size))
        arrays.sort(axis=2)
        _, stats = propagation._replicate_statistics(arrays)
        all_values = np.concatenate(arrays, axis=1)
        np.testing.assert_allclose(stats[:3], np.quantile(all_values, [.025, .5, .975], axis=1), atol=1e-14)
        np.testing.assert_allclose(stats[3], np.mean(all_values, axis=1), atol=1e-14)
        np.testing.assert_allclose(stats[4], np.std(all_values, axis=1), atol=1e-14)


def test_sampler_construction_is_serialized_but_requests_can_overlap(monkeypatch):
    original = propagation.qmc.Sobol
    guard, start = Lock(), Barrier(2)
    active = 0
    def checked_constructor(*args, **kwargs):
        nonlocal active
        with guard:
            active += 1
            assert active == 1, "Concurrent initialization of SciPy direction tables"
        try:
            time.sleep(.01)
            return original(*args, **kwargs)
        finally:
            with guard:
                active -= 1
    monkeypatch.setattr(propagation.qmc, "Sobol", checked_constructor)
    def solve(_):
        start.wait(timeout=10)
        return predict_isotopes(ForwardIsotopeConstraints(
            pco2_constraint=C("range", lower=100, upper=500)), surface=LinearSurface())
    with ThreadPoolExecutor(max_workers=2) as pool:
        results = list(pool.map(solve, range(2)))
    assert results[0] == results[1]


def test_cold_process_concurrent_sampling_matches_sequential():
    root = Path(__file__).resolve().parents[1]
    script = '''
import sys
sys.path.insert(0, 'code')
from concurrent.futures import ThreadPoolExecutor
from threading import Barrier
from forward_isotope_constraints import ForwardIsotopeConstraints, predict_isotopes
from updated_constrained_pco2_posterior import CoordinateConstraint
from updated_output_surface import load_updated_output_surface
load_updated_output_surface()
request = ForwardIsotopeConstraints(gpp_constraint=CoordinateConstraint('normal', center=290, sigma=29))
start = Barrier(2)
def solve(_):
    start.wait(timeout=10)
    return predict_isotopes(request)
with ThreadPoolExecutor(max_workers=2) as pool:
    results = list(pool.map(solve, range(2)))
assert results[0] == results[1] == predict_isotopes(request)
'''
    completed = subprocess.run([sys.executable, "-c", script], cwd=root,
                               capture_output=True, text=True, timeout=60)
    assert completed.returncode == 0, completed.stderr
    assert "Exception ignored" not in completed.stderr, completed.stderr


@pytest.mark.parametrize("inputs", [
    ForwardIsotopeConstraints(gpp_constraint=C("normal", center=290, sigma=29)),
    ForwardIsotopeConstraints(C("normal", center=30000, sigma=6000), C("normal", center=58, sigma=29), C("normal", center=.2, sigma=.2)),
    ForwardIsotopeConstraints(C("range", lower=50, upper=60000), C("range", lower=18.256264, upper=850), C("range", lower=.1, upper=2)),
])
def test_real_surface_propagation_converges(inputs):
    result = predict_isotopes(inputs)
    checks = result["numerical_checks"]
    assert checks["converged"]
    assert np.all(np.asarray(checks["replicate_error_permil"]) <= checks["tolerance_permil"])
    for iso in result["isotopes"].values():
        assert iso["interval95"][0] <= iso["median"] <= iso["interval95"][1]


@pytest.mark.parametrize("constraint", [C("fixed", center=49), C("range", lower=40, upper=100), C("normal", center=294, sigma=-1), C("fixed", center=float("nan"))])
def test_invalid_constraints_fail(constraint):
    with pytest.raises(ValueError):
        predict_isotopes(ForwardIsotopeConstraints(pco2_constraint=constraint))


def test_failed_convergence_returns_no_prediction(monkeypatch):
    monkeypatch.setattr(propagation, "_MAX_POWER", propagation._MIN_POWER)
    with pytest.raises(propagation.ForwardResolutionError):
        predict_isotopes(ForwardIsotopeConstraints(pco2_constraint=C("range", lower=100, upper=500)))


def test_api_and_workbook():
    payload = {"pco2_constraint": {"kind": "fixed", "center": 294},
               "gpp_constraint": {"kind": "fixed", "center": 290},
               "po2_constraint": {"kind": "fixed", "center": 1}}
    with TestClient(web_api.app) as client:
        response = client.post("/api/v1/forward/isotopes", json=payload)
        assert response.status_code == 200, response.text
        export = client.post("/api/v1/export/isotopes.xlsx", json=payload)
        assert export.status_code == 200
        workbook = load_workbook(BytesIO(export.content), read_only=True)
        assert workbook.sheetnames == ["Summary", "Input constraints", "Metadata"]
        rows = list(workbook["Summary"].values)
        assert any(row and row[0] == "delta18_conventional_permil: median" for row in rows)
        workbook.close()
        assert client.post("/api/v1/forward/isotopes", json={}).status_code == 422
        payload["po2_constraint"]["center"] = 3
        assert client.post("/api/v1/forward/isotopes", json=payload).status_code == 422
