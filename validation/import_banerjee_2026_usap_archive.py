"""Normalize the checksum-verified Banerjee et al. (2026) USAP workbook.

The raw workbook remains in ignored ``external_data``. This importer writes a
small, source-controlled CSV used by the validation audit and preserves the
authors' pristine-sample exclusions.
"""

from __future__ import annotations

import csv
import hashlib
from pathlib import Path

from openpyxl import load_workbook


ROOT = next(path for path in Path(__file__).resolve().parents if (path / ".project-root").exists())
ARCHIVE_PATH = ROOT / "external_data" / "ice_core_validation" / "banerjee_2026_usap_602070.xlsx"
OUTPUT_PATH = ROOT / "model_data" / "literature" / "banerjee_2026_supporting_tables.csv"
EXPECTED_MD5 = "2c0e4d4090f58cb790e098647141b219"
EXPECTED_SHA256 = "2bc30014d56169c37a9143ad263767398b6f520fc930dc1c2ac2dae1b90b35c6"
BANERJEE_INTERCEPT_PPM = 266.58
BANERJEE_SLOPE_MAGNITUDE = 1.6131
LEGACY_INTERCEPT_PPM = 256.9758
LEGACY_SLOPE_MAGNITUDE = 1.0862

FIELDS = (
    "table", "row", "core_id", "top_depth_m", "ar_age_ka",
    "excluded_from_pristine_co2_dataset", "measured_co2_ppm",
    "delta13c_permil", "delta15n_permil", "delta_o2_n2_permil",
    "delta_ar_n2_permil", "cap_delta17_rep1_ppm", "cap_delta17_rep2_ppm",
    "cap_delta17_average_ppm", "reconstructed_co2_ppm", "source_note",
)


def _digest(path: Path, algorithm: str) -> str:
    hasher = hashlib.new(algorithm)
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            hasher.update(block)
    return hasher.hexdigest()


def _value(value: object) -> object:
    if value is None or value in {"NaN", "N/A"}:
        return ""
    return value


def _exclusion(value: object) -> str:
    return "yes" if str(value).strip().upper() == "YES" else "no"


def _s1_rows(workbook: object) -> list[dict[str, object]]:
    sheet = workbook["ALHIC1901"]
    rows: list[dict[str, object]] = []
    for source_row, values in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
        if values[0] is None:
            continue
        note = ""
        if source_row == 6:
            note = "Archive replicate/average inconsistency: rep1 is 32.1524 ppm, rep2 is absent, average is 29.4164 ppm"
        elif source_row == 22:
            note = "Extreme altered sample excluded by Banerjee et al.; reconstructed CO2 is N/A"
        rows.append({
            "table": "S1",
            "row": source_row - 1,
            "core_id": values[0],
            "top_depth_m": _value(values[1]),
            "ar_age_ka": _value(values[2]),
            "excluded_from_pristine_co2_dataset": _exclusion(values[3]),
            "measured_co2_ppm": _value(values[4]),
            "delta13c_permil": _value(values[5]),
            "delta15n_permil": _value(values[6]),
            "delta_o2_n2_permil": _value(values[7]),
            "delta_ar_n2_permil": _value(values[8]),
            "cap_delta17_rep1_ppm": _value(values[9]),
            "cap_delta17_rep2_ppm": _value(values[10]),
            "cap_delta17_average_ppm": _value(values[11]),
            "reconstructed_co2_ppm": _value(values[12]),
            "source_note": note,
        })
    return rows


def _s2_rows(workbook: object) -> list[dict[str, object]]:
    sheet = workbook["ALHIC1502 1503"]
    rows: list[dict[str, object]] = []
    for values in sheet.iter_rows(min_row=2, values_only=True):
        if values[0] is None:
            continue
        rows.append({
            "table": "S2",
            "row": len(rows) + 1,
            "core_id": values[0],
            "top_depth_m": _value(values[1]),
            "ar_age_ka": _value(values[2]),
            "excluded_from_pristine_co2_dataset": _exclusion(values[3]),
            "measured_co2_ppm": _value(values[4]),
            "delta13c_permil": _value(values[5]),
            "delta15n_permil": _value(values[6]),
            "delta_o2_n2_permil": _value(values[7]),
            "delta_ar_n2_permil": "",
            "cap_delta17_rep1_ppm": "",
            "cap_delta17_rep2_ppm": "",
            "cap_delta17_average_ppm": _value(values[8]),
            "reconstructed_co2_ppm": _value(values[9]),
            "source_note": "",
        })
    return rows


def run(archive_path: Path = ARCHIVE_PATH, output_path: Path = OUTPUT_PATH) -> list[dict[str, object]]:
    md5 = _digest(archive_path, "md5")
    sha256 = _digest(archive_path, "sha256")
    if md5 != EXPECTED_MD5 or sha256 != EXPECTED_SHA256:
        raise ValueError(f"Banerjee archive checksum mismatch: md5={md5}, sha256={sha256}")

    workbook = load_workbook(archive_path, data_only=True, read_only=True)
    if workbook["README"]["C19"].value.strip() != "266.58 - 1.6131*Delta17O-O2,grav":
        # The actual cell uses a Greek Delta; normalize it without weakening the check.
        equation = workbook["README"]["C19"].value.strip().replace("Delta", "Δ")
        if equation != "266.58 - 1.6131*Δ17O-O2,grav":
            raise ValueError(f"Unexpected archived reconstruction equation: {equation}")

    rows = _s1_rows(workbook) + _s2_rows(workbook)
    for row in rows:
        delta = row["cap_delta17_average_ppm"]
        reported = row["reconstructed_co2_ppm"]
        if delta == "" or reported == "":
            continue
        if row["table"] == "S1":
            predicted = BANERJEE_INTERCEPT_PPM - BANERJEE_SLOPE_MAGNITUDE * float(delta)
        else:
            predicted = LEGACY_INTERCEPT_PPM - LEGACY_SLOPE_MAGNITUDE * float(delta)
        tolerance_ppm = 0.002 if row["table"] == "S1" else 1e-9
        if abs(predicted - float(reported)) > tolerance_ppm:
            raise ValueError(f"Archived reconstruction mismatch in {row['table']} row {row['row']}")

    output_path.parent.mkdir(parents=True, exist_ok=True)
    with output_path.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, fieldnames=FIELDS)
        writer.writeheader()
        writer.writerows(rows)
    return rows


if __name__ == "__main__":
    imported = run()
    print(f"Imported {len(imported)} checksum-verified rows to {OUTPUT_PATH}")
