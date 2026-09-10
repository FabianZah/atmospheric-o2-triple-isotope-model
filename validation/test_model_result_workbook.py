"""Streaming export preserves scientific data, formatting and temporary-file isolation."""

from copy import deepcopy
from io import BytesIO
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.worksheet import _writer
import pytest

import model_result_workbook as exporter


@pytest.fixture
def envelope():
    return {
        "calculation": "coordinate_inference",
        "result": {
            "inputs": {
                "target_air_cap_delta17_permil": -8.0,
                "measurement_sigma_permil": 0.015,
                "credible_mass": 0.95,
                "constraints": {"pO2": {"kind": "fixed", "center": 0.2}},
            },
            "solve_for": "pCO2",
            "solve_axis": [50.0, 100.0, 1000.0],
            "solve_marginal_probability_mass": [0.0, 0.3, 0.7],
            "solve_marginal_density": [0.0, 0.003, 0.007],
            "posterior_median": 800.1234567890123,
            "equal_tailed_credible_interval": [55.5555555555555, 999.1234567890123],
            "solve_boundary_sensitive": False,
            "probability_scope": "conditional",
            "effective_constraint_bounds": {"pO2": [0.2, 0.2]},
            "field_x_coordinate": "pCO2",
            "field_y_coordinate": "GPP",
            "field_x_axis": [50.0, 100.0, 1000.0],
            "field_y_axis": [29.0, 290.0, 580.0],
            "field_probability_mass": [0.0, 0.0, 0.0, 0.0, 1e-250, 0.3, 0.0, 0.0, 0.7],
            "field_density": [0.0, 0.0, 0.0, 0.0, 1e-245, 3.0, 0.0, 0.0, 7.0],
            "field_hpd_mask": [False, False, False, False, False, True, False, False, True],
            "field_hpd_density_threshold": 3.0,
            "field_hpd_probability_mass": 1.0,
        },
    }


@pytest.mark.parametrize("coordinates", [("pCO2", "GPP"), ("pCO2", "pO2"), ("GPP", "pO2")])
def test_streamed_field_preserves_values_types_axes_and_layout(envelope, coordinates):
    result = envelope["result"]
    result["field_x_coordinate"], result["field_y_coordinate"] = coordinates
    before = deepcopy(envelope)
    content = exporter.build_coordinate_inference_workbook(envelope, {"isotope_source": "Direct air O2"})
    assert envelope == before
    workbook = load_workbook(BytesIO(content), data_only=False)
    assert workbook.sheetnames == ["Summary", "Posterior", "Joint probability"]
    field = workbook["Joint probability"]
    rows = list(field.iter_rows(min_row=4, values_only=True))
    expected = []
    for i, x in enumerate(result["field_x_axis"]):
        for j, y in enumerate(result["field_y_axis"]):
            k = i * 3 + j
            mass, density = result["field_probability_mass"][k], result["field_density"][k]
            if i and j and mass == 0 and density == 0:
                continue
            expected.append((x, exporter._unit(coordinates[0]), y, exporter._unit(coordinates[1]),
                             mass, density, result["field_hpd_mask"][k]))
    assert rows == expected
    assert rows[4][4] == 1e-250  # Positive tails are retained, however small.
    assert sum(row[4] for row in rows) == 1.0
    assert all(type(row[6]) is bool for row in rows)
    assert set(row[0] for row in rows) == set(result["field_x_axis"])
    assert set(row[2] for row in rows) == set(result["field_y_axis"])
    for sheet, columns in [(workbook["Summary"], "C"), (workbook["Posterior"], "D"), (field, "G")]:
        assert sheet.freeze_panes == "A4"
        assert str(sheet.merged_cells) == f"A1:{columns}1"
        assert sheet.auto_filter.ref == f"A3:{columns}{sheet.max_row}"
        assert sheet.row_dimensions[1].height == 24
        assert sheet["A1"].fill.fgColor.rgb == "00123238"
        assert sheet["A3"].fill.fgColor.rgb == "00006F71"
        assert sheet["A4"].alignment.wrap_text
    assert field.column_dimensions["E"].width == 22
    assert field["A4"].number_format == "0.0000000000"
    assert field["E4"].number_format == "0.0000000000E+00"
    summary = {row[0]: row[1] for row in workbook["Summary"].iter_rows(min_row=4, values_only=True)}
    assert summary["posterior_median"] == pytest.approx(result["posterior_median"], rel=1e-15)
    assert summary["boundary_sensitive"] is False
    assert summary["target_air_delta18O_VSMOW"] is None


def test_streamed_table_retains_no_cell_grid_and_treats_metadata_as_text():
    with exporter._streaming_workbook() as workbook:
        assert workbook.write_only
        table = exporter._StreamingTable(workbook, "Data", title="Data", headers=("Value", "Flag"), widths=(30, 12))
        text = '=HYPERLINK("https://example.invalid", "literal")'
        for _ in range(1000):
            table.append((text, False))
        assert not getattr(table.sheet, "_cells", {})
        table.finish()
        buffer = BytesIO()
        workbook.save(buffer)
    restored = load_workbook(BytesIO(buffer.getvalue()), read_only=True, data_only=False)
    rows = list(restored.active.iter_rows(min_row=4))
    assert len(rows) == 1000
    assert all(row[0].value == text and row[0].data_type == "s" and row[1].value is False for row in rows)
    restored.close()


@pytest.mark.parametrize("failure", ["append", "save", None])
def test_export_temp_files_are_removed_without_touching_another_workbook(envelope, monkeypatch, failure):
    with exporter._streaming_workbook() as other:
        table = exporter._StreamingTable(other, "Other", title="Other", headers=("A", "B"), widths=(10, 10))
        table.append((1, 2))
        other_file = Path(table.sheet._writer.out)
        initial_files = set(_writer.ALL_TEMP_FILES)
        if failure == "append":
            original = exporter._StreamingTable.append

            def fail_after_row(self, values):
                original(self, values)
                raise RuntimeError("injected append failure")

            monkeypatch.setattr(exporter._StreamingTable, "append", fail_after_row)
        elif failure == "save":
            def fail_save(*args, **kwargs):
                raise RuntimeError("injected save failure")

            monkeypatch.setattr(exporter.Workbook, "save", fail_save)
        if failure:
            with pytest.raises(RuntimeError, match=f"injected {failure} failure"):
                exporter.build_coordinate_inference_workbook(envelope, {"isotope_source": "Direct air O2"})
        else:
            exporter.build_coordinate_inference_workbook(envelope, {"isotope_source": "Direct air O2"})
        assert set(_writer.ALL_TEMP_FILES) == initial_files
        assert other_file.is_file()
    assert not other_file.exists()


def test_streaming_requires_incremental_xml_backend_before_creating_files(monkeypatch):
    initial_files = set(_writer.ALL_TEMP_FILES)
    monkeypatch.setattr(exporter, "LXML", False)
    with pytest.raises(RuntimeError, match="Streaming XLSX export requires lxml"):
        with exporter._streaming_workbook():
            pytest.fail("buffered XML backend must not be used")
    assert set(_writer.ALL_TEMP_FILES) == initial_files
