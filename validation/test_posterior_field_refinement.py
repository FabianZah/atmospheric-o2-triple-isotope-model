"""Contracts for conditional marginal-field refinement, without large model runs."""
from dataclasses import fields
from types import SimpleNamespace

import numpy as np
import pytest
from scipy.ndimage import label

import posterior_field_refinement as refinement
from updated_constrained_pco2_posterior import (
    ConstrainedCoordinateInput, ConstrainedCoordinateResult, _pair_hpd, _trapezoid_weights,
)


@pytest.fixture(autouse=True)
def clear_refinement_cache(monkeypatch):
    monkeypatch.setattr(refinement, "_LAST_REFINEMENT", None)


def test_wide_connected_field_stays_on_fast_path():
    mask = np.zeros((25,25), dtype=bool)
    mask[5:20,5:20] = True
    mass = mask.astype(float)/mask.sum()
    assert not refinement.field_resolution_trigger(mask, mass)["needed"]


def test_islands_and_thin_ridges_request_recomputation():
    mask = np.eye(25, dtype=bool)
    assert refinement.field_resolution_trigger(mask, mask/25)["needed"]
    mask = np.zeros((25,25), dtype=bool)
    mask[3:9,3:9] = mask[16:22,16:22] = True
    check = refinement.field_resolution_trigger(mask, mask/mask.sum())
    assert check["needed"] and check["components"] == 2


def test_focused_axis_retains_outer_domain_and_real_boundary_support():
    axis = np.geomspace(50,60000,181)
    mass = np.exp(-.5*((np.log(axis)-np.log(5000))/.1)**2)
    refined = refinement._refined_axis(axis, mass, mass > mass.max()*.1, "pCO2")
    assert refined[0] == 50 and refined[-1] == 60000
    assert np.all(np.diff(refined) > 0)
    assert len(refined[(refined>4000)&(refined<6000)]) > 100
    boundary_mass = np.exp(-.5*((axis-60000)/5000)**2)
    refined = refinement._refined_axis(axis, boundary_mass, boundary_mass > .1, "pCO2")
    assert refined[-1] == 60000


def _toy_result(solve_for="pCO2"):
    x,y,z = np.linspace(1,10,25),np.linspace(1,10,25),np.linspace(.1,2,17)
    mass = np.eye(25)/25
    density,mask,hpd = _pair_hpd(mass,x,y,.95)
    values = {field.name: None for field in fields(ConstrainedCoordinateResult)}
    values.update(inputs=ConstrainedCoordinateInput(solve_for,0.,1.,{}),
                  solve_for=solve_for, field_x_coordinate="pCO2",field_y_coordinate="GPP",
                  field_shape=mass.shape,field_probability_mass=mass.ravel(),field_density=density.ravel(),
                  field_hpd_mask=mask.ravel(),field_x_axis=x,field_y_axis=y,
                  field_hpd_probability_mass=hpd,field_hpd_density_threshold=float(density[mask].min()),
                  posterior_median=5.,equal_tailed_credible_interval=(2.,8.),
                  constraint_posterior_medians={c: 1. for c in ("pCO2","GPP","pO2") if c!=solve_for},
                  constraint_equal_tailed_credible_intervals={},effective_grid_sizes={"pCO2":25,"GPP":25,"pO2":17},
                  probability_scope="Test model.")
    nuisance="pO2"
    if solve_for=="pO2":
        values["field_y_coordinate"]="pO2"
        nuisance="GPP"
    request=SimpleNamespace(integrate_coordinate=nuisance,pco2_prior="uniform",gpp_prior="uniform",po2_prior="uniform")
    return ConstrainedCoordinateResult(**values),SimpleNamespace(inputs=request,axes={nuisance:z})


@pytest.mark.parametrize("solve_for", ("pCO2","GPP","pO2"))
def test_streamed_marginals_preserve_real_separate_regions(monkeypatch, solve_for):
    monkeypatch.setattr(refinement,"PAIR_AXIS_FLOORS",{"pCO2":101,"GPP":101,"pO2":101})
    monkeypatch.setattr(refinement,"CHUNK_VOLUME_POINTS",4000)
    def likelihood(surface, request, point, z, nuisance):
        other="GPP" if nuisance=="pO2" else "pO2"
        x,y,z=point["pCO2"],point[other],point[nuisance]
        a=-.5*((x-3)/.25)**2-.5*((y-3)/.25)**2
        b=-.5*((x-8)/.25)**2-.5*((y-8)/.25)**2
        return np.logaddexp(a,b)-.5*((z-1)/.2)**2, {
            "forward_points":x.size,"maximum_level":1,
            "maximum_final_relative_integral_change":0.,"maximum_final_response_error_sigma":0.}
    monkeypatch.setattr(refinement,"integrated_coordinate_log_likelihood",likelihood)
    result,joint=_toy_result(solve_for)
    updated=refinement.refine_integrated_pair_field(result,joint,None)
    mass=np.asarray(updated.field_probability_mass).reshape(updated.field_shape)
    mask=np.asarray(updated.field_hpd_mask).reshape(updated.field_shape)
    density=np.asarray(updated.field_density).reshape(updated.field_shape)
    assert label(mask,np.ones((3,3)))[1] == 2
    assert mass.sum() == pytest.approx(1.)
    assert updated.field_hpd_probability_mass == pytest.approx(.95,abs=.003)
    assert np.array_equal(mask,density>=updated.field_hpd_density_threshold)
    expected=mass.sum(axis=1 if solve_for=="pCO2" else 0)
    assert np.asarray(updated.solve_marginal_probability_mass) == pytest.approx(expected)
    assert updated.constraint_posterior_medians[joint.inputs.integrate_coordinate] == pytest.approx(.94,abs=.015)
    assert updated.field_refinement_diagnostics["refined_components"] == 2
    assert updated.field_x_axis[0] == 1 and updated.field_x_axis[-1] == 10
    monkeypatch.setattr(refinement, "integrated_coordinate_log_likelihood",
                        lambda *args: pytest.fail("identical export request recalculated the dense field"))
    assert refinement.refine_integrated_pair_field(result,joint,None) is updated
    if solve_for=="pCO2":
        from io import BytesIO
        from openpyxl import load_workbook
        from model_result_workbook import build_coordinate_inference_workbook

        payload=build_coordinate_inference_workbook(
            {"calculation":"constrained_coordinate_posterior","result":updated.as_dict()},
            {"isotope_source":"Direct air"})
        workbook=load_workbook(BytesIO(payload),read_only=True)
        summary={row[0]:row[1] for row in workbook["Summary"].iter_rows(min_row=4,values_only=True)}
        assert summary["field_refinement.status"] == "refined"
        assert summary["posterior_median"] == pytest.approx(updated.posterior_median)
        rows=list(workbook["Joint probability"].iter_rows(min_row=4,values_only=True))
        assert sum(row[4] for row in rows) == pytest.approx(1.)
        assert sum(row[4] for row in rows if row[6]) == pytest.approx(updated.field_hpd_probability_mass)
        workbook.close()


def test_unneeded_refinement_does_not_evaluate_model(monkeypatch):
    result,joint=_toy_result()
    monkeypatch.setattr(refinement,"field_resolution_trigger",lambda mask,mass:{"needed":False})
    assert refinement.refine_integrated_pair_field(result,joint,None) is result


def test_calculation_budget_fails_explicitly(monkeypatch):
    result,joint=_toy_result()
    monkeypatch.setattr(refinement,"MAX_FIELD_POINTS",10)
    with pytest.raises(RuntimeError,match="grid budget"):
        refinement.refine_integrated_pair_field(result,joint,None)


def test_general_sulfate_path_is_not_replaced_with_gaussian_air():
    result,joint=_toy_result()
    joint.inputs.integrate_coordinate=None
    assert refinement.refine_integrated_pair_field(result,joint,None) is result
