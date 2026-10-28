"""Both biological configurations cross the same public API and export boundary."""
from io import BytesIO

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from web_api import app


@pytest.mark.parametrize("name", ["modern_reference", "marine_endmember"])
def test_configuration_forward_cache_and_workbook(name):
    client = TestClient(app)
    payload = {"biosphere_configuration":name,
        "pco2_constraint":{"kind":"fixed","center":294.},
        "gpp_constraint":{"kind":"fixed","center":290.},
        "po2_constraint":{"kind":"fixed","center":1.}}
    result = client.post("/api/v1/forward/isotopes",json=payload)
    assert result.status_code == 200, result.text
    result = result.json()
    assert result["result"]["inputs"]["biosphere_configuration"] == name
    assert result["provenance"]["biosphere_configuration"] == name
    response = client.post("/api/v1/export/isotopes.xlsx",json=payload)
    assert response.status_code == 200, response.text
    book = load_workbook(BytesIO(response.content),read_only=True,data_only=True)
    summary = {row[0]:row[1] for row in book["Summary"].iter_rows(min_row=4,values_only=True)}
    assert summary["biological_configuration"] == name
    assert summary["software_version"] == "0.2.0"
    book.close()


def test_alternating_requests_do_not_reuse_other_biosphere_cache():
    client = TestClient(app)
    values = []
    for name in ["modern_reference", "marine_endmember", "modern_reference"]:
        response = client.post("/api/v1/forward",json={"biosphere_configuration":name})
        assert response.status_code == 200, response.text
        body = response.json()
        values.append(body["result"]["central_cap_delta17_prime_permil"])
        assert body["provenance"]["biosphere_configuration"] == name
    assert values[0] == values[2]
    assert abs(values[0]-values[1]) > .01


@pytest.mark.parametrize("route,extra", [
    ("state-step",{"final":{"biosphere_configuration":"marine_endmember"}}),
    ("photosynthesis-step",{"photosynthesis_fraction":1.}),
    ("pco2-trajectory",{"final_pco2_ppm":294.,"transition_duration_years":100.}),
])
def test_marine_time_response_route_keeps_configured_initial_state(route, extra):
    client = TestClient(app)
    initial = {"biosphere_configuration":"marine_endmember"}
    request = {"initial":initial,"duration_years":100.,"sample_count":5,
               "equilibrium_search_max_years":100.,**extra}
    response = client.post("/api/v1/transients/"+route,json=request)
    assert response.status_code == 200, response.text
    result = response.json()
    assert result["provenance"]["biosphere_configuration"] == "marine_endmember"
    values = [s["cap_delta17_prime_permil"] for s in result["result"]["states"]]
    assert max(values)-min(values) < 1e-6


def test_marine_joint_and_conditional_probability_routes():
    client = TestClient(app)
    point = {"biosphere_configuration":"marine_endmember", "p_o2_pal":.7,
             "p_co2_ppm":5000.,"gpp_pgC_per_year":180.}
    target = client.post("/api/v1/forward",json=point).json()["result"]["central_cap_delta17_prime_permil"]
    base = {**point,"target_air_cap_delta17_permil":target,"measurement_sigma_permil":.06}
    for route, extra in [("conditional",{"solve_for":"pCO2","grid_size":257}),
                         ("joint",{"free_coordinates":["pCO2","GPP"],"pco2_grid_size":41,"gpp_grid_size":41})]:
        response = client.post("/api/v1/posterior/"+route,json={**base,**extra})
        assert response.status_code == 200, response.text
        result = response.json()
        assert result["provenance"]["biosphere_configuration"] == "marine_endmember"
        assert result["result"]["inputs"]["biosphere_configuration"] == "marine_endmember"


def test_unknown_configuration_is_rejected():
    response = TestClient(app).post("/api/v1/forward",json={"biosphere_configuration":"marine_unverified"})
    assert response.status_code == 422
