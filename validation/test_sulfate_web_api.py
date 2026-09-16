"""Sulfate browser contracts, calculation limits and reproducible XLSX exports."""

from dataclasses import asdict
from io import BytesIO
import math

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from web_api import app
from test_sulfate_uncertainty import synthetic_observation, normal_nonair_browser_observation
from sulfate_uncertainty import (
    SulfateComputeLimitError, exact_sulfate_likelihood, sulfate_computation_budget,
)


def payload(coordinate="pCO2"):
    observation = asdict(synthetic_observation())
    observation["primary_signal_assumed"] = True
    result = {"solve_for": coordinate, "sulfate": observation}
    for name, field, value in (("pCO2", "pco2_constraint", 10000.),
                               ("GPP", "gpp_constraint", 72.5), ("pO2", "po2_constraint", 0.5)):
        if name != coordinate:
            result[field] = {"kind": "fixed", "center": value}
    return result


@pytest.mark.parametrize("coordinate,truth", [("pCO2", 10000.), ("GPP", 72.5), ("pO2", 0.5)])
def test_sulfate_browser_solves_each_coordinate_without_a_fabricated_air_target(coordinate, truth):
    with TestClient(app) as client:
        response = client.post("/api/v1/inference/coordinate", json=payload(coordinate))
    assert response.status_code == 200, response.text
    result = response.json()["result"]
    low, high = result["equal_tailed_credible_interval"]
    assert low < truth < high
    assert result["inputs"]["target_air_cap_delta17_permil"] is None
    assert result["sulfate_likelihood_diagnostics"]["status"] == "converged"
    assert "preservation is assumed, not assessed" in result["inputs"]["sulfate"]["assumption_note"]


@pytest.mark.parametrize("incorporation,background", [
    ({"kind": "range", "lower": .15, "upper": .25}, {"kind": "fixed", "center": -.02}),
    ({"kind": "fixed", "center": .2}, {"kind": "normal", "center": -.02, "sigma": .03}),
    ({"kind": "normal", "center": .2, "sigma": .02}, {"kind": "range", "lower": -.04, "upper": 0.}),
])
def test_browser_process_constraint_modes(incorporation, background):
    from web_api import ConstrainedCoordinateRequest
    request = payload()
    request["sulfate"]["incorporation"] = incorporation
    request["sulfate"]["background"] = background
    parsed = ConstrainedCoordinateRequest.model_validate(request).solver_input()
    assert parsed.sulfate.incorporation.kind == incorporation["kind"]
    assert parsed.sulfate.background.kind == background["kind"]


def test_browser_nonair_normal_uncertainty_converges_and_exports_diagnostics():
    request = {
        "solve_for": "pCO2", "sulfate": asdict(normal_nonair_browser_observation()),
        "gpp_constraint": {"kind": "fixed", "center": 290.0},
        "po2_constraint": {"kind": "fixed", "center": 1.0},
    }
    with TestClient(app) as client:
        response = client.post("/api/v1/inference/coordinate", json=request)
    assert response.status_code == 200, response.text
    result = response.json()["result"]
    diagnostics = result["sulfate_likelihood_diagnostics"]
    assert diagnostics["status"] == "converged"
    assert diagnostics["maximum_inner_compatibility_error_estimate"] >= 0.
    assert result["posterior_median"] == pytest.approx(1715.15, abs=0.1)
    assert result["equal_tailed_credible_interval"] == pytest.approx((844.42, 2860.07), abs=0.1)
    assert result["inputs"]["sulfate"]["background"]["sigma"] == 0.05
    assert diagnostics["settings"]["relative_tolerance"] <= 2e-4


@pytest.mark.parametrize("change", [
    {"primary_signal_assumed": False}, {"assumption_note": " "},
    {"incorporation": {"kind": "fixed", "center": 20.}},
    {"incorporation": {"kind": "fixed", "center": 1.}},
    {"background": {"kind": "range", "lower": .1, "upper": -.1}},
    {"isotope_error_correlation": 1.}, {"theta_air_to_sulfate": 0.},
    {"delta18_sigma_permil": -1.}, {"cap_delta17_sigma_permil": 0.},
])
def test_sulfate_invalid_or_incompatible_process_is_rejected(change):
    request = payload()
    request["sulfate"].update(change)
    with TestClient(app) as client:
        response = client.post("/api/v1/inference/coordinate", json=request)
    assert response.status_code == 422


def test_simplified_form_needs_no_confirmation_or_free_text_and_preserves_results():
    from web_api import ConstrainedCoordinateRequest
    request = payload()
    explicit = ConstrainedCoordinateRequest.model_validate(request).solver_input()
    request["sulfate"].pop("assumption_note")
    request["sulfate"].pop("primary_signal_assumed")
    simplified = ConstrainedCoordinateRequest.model_validate(request).solver_input()
    explicit_values = asdict(explicit.sulfate)
    simplified_values = asdict(simplified.sulfate)
    explicit_values.pop("assumption_note")
    note = simplified_values.pop("assumption_note")
    assert simplified_values == explicit_values
    assert note == "Conditional primary-sulfate transfer; preservation is assumed, not assessed."
    with TestClient(app) as client:
        response = client.post("/api/v1/inference/coordinate", json=request)
    assert response.status_code == 200, response.text
    result = response.json()["result"]
    assert result["inputs"]["sulfate"]["assumption_note"] == note
    low, high = result["equal_tailed_credible_interval"]
    assert low < 10000. < high


def test_air_and_sulfate_inputs_are_mutually_exclusive():
    request = payload()
    request.update(target_air_cap_delta17_permil=-10., measurement_sigma_permil=.1)
    with TestClient(app) as client:
        response = client.post("/api/v1/inference/coordinate", json=request)
    assert response.status_code == 422
    assert "mutually exclusive" in response.text


def test_budget_accumulates_across_calls_and_resets_after_failure():
    observation = synthetic_observation()
    with pytest.raises(SulfateComputeLimitError):
        with sulfate_computation_budget(max_roots=20):
            exact_sulfate_likelihood(-12.875, 25.9, observation)
            exact_sulfate_likelihood(-12.875, 25.9, observation)
    assert exact_sulfate_likelihood(-12.875, 25.9, observation).diagnostics["status"] == "converged"


def test_browser_budget_failure_returns_an_error_not_a_partial_posterior(monkeypatch):
    import web_api
    monkeypatch.setattr(web_api, "sulfate_computation_budget", lambda: sulfate_computation_budget(max_roots=1))
    with TestClient(app) as client:
        response = client.post("/api/v1/inference/coordinate", json=payload())
    assert response.status_code == 503
    assert response.json()["code"] == "sulfate_calculation_limit"
    assert "computational limit" in response.json()["detail"]
    assert "result" not in response.json()


def test_browser_sulfate_budget_allows_the_audited_multi_constraint_work(monkeypatch):
    import sulfate_uncertainty as module
    monkeypatch.setattr(module,"monotonic",lambda:0.)
    with sulfate_computation_budget():
        module._COMPUTE_BUDGET.get().claim(410_201_767)
        monkeypatch.setattr(module,"monotonic",lambda:550.)
    assert module._COMPUTE_BUDGET.get() is None
    with pytest.raises(SulfateComputeLimitError):
        with sulfate_computation_budget():
            module._COMPUTE_BUDGET.get().claim(1_000_000_001)


def test_time_budget_is_checked_before_return_even_without_another_root(monkeypatch):
    import sulfate_uncertainty
    monkeypatch.setattr(sulfate_uncertainty, "monotonic", lambda: 0.)
    with pytest.raises(SulfateComputeLimitError):
        with sulfate_computation_budget(max_seconds=1.):
            monkeypatch.setattr(sulfate_uncertainty, "monotonic", lambda: 2.)
    assert sulfate_uncertainty._COMPUTE_BUDGET.get() is None


def test_sulfate_workbook_contains_process_metadata_and_no_false_air_measurement():
    request = payload()
    request["sulfate"]["assumption_note"] = '=HYPERLINK("https://example.invalid", "Not a formula")'
    with TestClient(app) as client:
        response = client.post("/api/v1/export/coordinate.xlsx", json={
            "inference": request, "context": {"isotope_source": "Sulfate"}})
    assert response.status_code == 200, response.text
    workbook = load_workbook(BytesIO(response.content), data_only=False)
    assert "Sulfate transfer" in workbook.sheetnames
    summary = {row[0]: row[1] for row in workbook["Summary"].iter_rows(min_row=4, values_only=True)}
    assert summary["isotope_source"] == "Sulfate"
    assert "target_air_Delta_prime_17O_0.528" not in summary
    assert summary["sulfate_delta18O_VSMOW"] == 15.
    metadata = {row[0]: row[1] for row in workbook["Sulfate transfer"].iter_rows(min_row=4, values_only=True)}
    assert metadata["observation_and_process.incorporation.center"] == .2
    assert metadata["observation_and_process.background.center"] == -.02
    assert metadata["observation_and_process.isotope_error_correlation"] == 0.
    assert metadata["likelihood_diagnostics.status"] == "converged"
    assert metadata["observation_and_process.assumption_note"].endswith(request["sulfate"]["assumption_note"])
    assert "confirmed" not in metadata["observation_and_process.assumption_note"]
    assert all(cell.data_type != "f" for sheet in workbook for row in sheet for cell in row)


def test_sulfate_workbook_source_cannot_be_mislabelled():
    with TestClient(app) as client:
        response = client.post("/api/v1/export/coordinate.xlsx", json={
            "inference": payload(), "context": {"isotope_source": "Direct air O2"}})
    assert response.status_code == 422


def test_sulfate_controls_and_help_are_shipped_with_the_app():
    with TestClient(app) as client:
        html = client.get("/").text
    for expected in ('data-source="sulfate"', 'id="sulfate-inputs"', 'id="sulfate-f-mode"',
                     'id="sulfate-b-mode"', 'id="sulfate-advanced"', 'id="sulfate-fractionation"',
                     "What does the sulfate input require?"):
        assert expected in html
    assert 'id="sulfate-primary"' not in html
    assert 'id="sulfate-assumptions"' not in html
    assert 'id="sulfate-correlation"' not in html
    assert "independent Gaussian errors" in html
    assert "Non-air oxygen component" in html
    assert "The prefilled isotope values and analytical errors are illustrative examples." in html
    assert 'id="sulfate-advanced" class="sulfate-advanced" open' not in html
    for field, value in (("sulfate-d17", "-0.200"), ("sulfate-d18", "15.000"),
                         ("sulfate-d17-sigma", "0.030"), ("sulfate-d18-sigma", "0.500")):
        assert f'id="{field}" type="text" inputmode="decimal" value="{value}" placeholder="Insert value"' in html
    assert 'id="sulfate-log18"' in html
    assert 'id="sulfate-shift17"' in html
    assert 'id="sulfate-alpha"' not in html
    assert 'value="peng_2026_irreversible"' in html
    assert 'value="peng_2026_equilibrium"' in html


def test_sulfate_range_labels_and_concise_help():
    with TestClient(app) as client:
        html = client.get("/").text
    for bound in ("lower", "upper"):
        assert f'for="sulfate-b-{bound}">Δ′<sup>17</sup>O {bound} <span>‰</span>' in html
    guide = html.split('id="view-guide"', 1)[1].split('id="view-references"', 1)[0]
    for question in (
        "Why is there no sulfate non-air δ<sup>18</sup>O input?",
        "Which sulfate O<sub>2</sub> incorporation treatment should I use?",
        "What do the sulfate custom isotope shifts mean?",
    ):
        assert f"<summary>{question}</summary>" in guide
    assert "Bao et al. (2008) established" not in guide
    assert "What uncertainty is included for sulfate?" not in guide
    assert "preserved or independently reconstructed primary sulfate signal" in guide
    assert "any sulfate-process constraints entered by the user" in guide
    assert "fixed 25% contribution" in guide
    assert html.index('id="references-proxy"') < html.index('id="references-sulfate"')
    assert html.index('id="references-sulfate"') < html.index('id="references-domain"')
    assert html.count('href="https://doi.org/10.1038/nature06959"') == 1


def named_payload(treatment="peng_2026_irreversible"):
    request = payload()
    s = request["sulfate"]
    for field in ("alpha18_air_to_sulfate", "alpha17_air_to_sulfate", "theta_air_to_sulfate"):
        s.pop(field, None)
    s["fractionation_treatment"] = treatment
    s["incorporation"] = {"kind": "fixed", "center": .25}
    return request


@pytest.mark.parametrize("treatment,shift", [("peng_2026_irreversible", .2916),
                                          ("peng_2026_equilibrium", .02835), ("none", 0.)])
def test_named_treatments_use_server_coefficients_and_record_provenance(treatment, shift):
    from web_api import ConstrainedCoordinateRequest
    from sulfate_to_air import SulfateProcessAssumptions, exact_sulfate_from_air
    from updated_output_surface import load_updated_output_surface, UpdatedOutputSurfaceInput
    request = named_payload(treatment)
    parsed = ConstrainedCoordinateRequest.model_validate(request).solver_input().sulfate
    transfer = parsed.fractionation_metadata()
    assert transfer["anomaly_shift_528_permil"] == pytest.approx(shift, abs=1e-12)
    if treatment != "none":
        assert transfer["source_doi"] == "10.1016/j.gca.2025.11.036"
        assert transfer["temperature_c"] == 25.
        assert transfer["required_incorporation_fraction"] == .25
    # Make an independently scalar-predicted sulfate at a known atmosphere.
    surface = load_updated_output_surface()
    state = surface.evaluate(UpdatedOutputSurfaceInput(p_o2_pal=.5, p_co2_ppm=10000., gpp_pgC_per_year=72.5))
    process = SulfateProcessAssumptions(parsed.background.center, parsed.alpha18_air_to_sulfate,
        parsed.theta_air_to_sulfate, "Synthetic pathway", parsed.alpha17_air_to_sulfate)
    request["sulfate"]["measured_cap_delta17_permil"] = exact_sulfate_from_air(
        air_cap_delta17_permil=state.central_cap_delta17_prime_permil,
        air_delta18_permil=1000.*math.expm1(state.central_delta18_prime_permil/1000.),
        sulfate_delta18_permil=15., fraction=.25, process=process)
    with TestClient(app) as client:
        response = client.post("/api/v1/inference/coordinate", json=request)
    assert response.status_code == 200, response.text
    result = response.json()["result"]
    lo, hi = result["equal_tailed_credible_interval"]
    assert lo < 10000. < hi
    assert result["sulfate_likelihood_diagnostics"]["fractionation"] == transfer


@pytest.mark.parametrize("change", [
    {"incorporation": {"kind": "range", "lower": .2, "upper": .3}},
    {"incorporation": {"kind": "fixed", "center": .3}},
    {"alpha18_air_to_sulfate": 1., "theta_air_to_sulfate": .528},
    {"alpha17_air_to_sulfate": .99},
    {"fractionation_treatment": "unknown"},
])
def test_named_treatment_cannot_silently_change_coefficients_or_stoichiometry(change):
    request = named_payload()
    request["sulfate"].update(change)
    with TestClient(app) as client:
        response = client.post("/api/v1/inference/coordinate", json=request)
    assert response.status_code == 422


def test_named_treatment_workbook_records_both_factors_shifts_and_source():
    with TestClient(app) as client:
        response = client.post("/api/v1/export/coordinate.xlsx", json={
            "inference": named_payload(), "context": {"isotope_source": "Sulfate"}})
    assert response.status_code == 200, response.text
    workbook = load_workbook(BytesIO(response.content), data_only=True)
    meta = {row[0]: row[1] for row in workbook["Sulfate transfer"].iter_rows(min_row=4, values_only=True)}
    assert meta["observation_and_process.fractionation_treatment"] == "peng_2026_irreversible"
    assert meta["observation_and_process.alpha17_air_to_sulfate"] == pytest.approx(math.exp(-.0216*.5145))
    assert meta["likelihood_diagnostics.fractionation.source_doi"] == "10.1016/j.gca.2025.11.036"
    assert meta["likelihood_diagnostics.fractionation.anomaly_shift_528_permil"] == pytest.approx(.2916)


def test_custom_zero_18_shift_is_not_treated_as_no_fractionation():
    from dataclasses import replace
    from web_api import ConstrainedCoordinateRequest
    request = payload()
    request["sulfate"].update(alpha18_air_to_sulfate=1., theta_air_to_sulfate=None,
                              alpha17_air_to_sulfate=math.exp(-.02376/1000))
    observation = ConstrainedCoordinateRequest.model_validate(request).solver_input().sulfate
    shifted = exact_sulfate_likelihood(-12.875, 25.9, observation)
    unchanged = exact_sulfate_likelihood(-12.875, 25.9, replace(observation, alpha17_air_to_sulfate=1.))
    assert shifted.log_likelihood != pytest.approx(unchanged.log_likelihood, abs=1e-6)
    assert shifted.diagnostics["fractionation"]["anomaly_shift_528_permil"] == pytest.approx(-.02376)
